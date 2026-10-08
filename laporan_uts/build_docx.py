# -*- coding: utf-8 -*-
"""Membangun Laporan UTS TIK Network Explorer (DOCX) sesuai template PDF 1."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

BASE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(BASE, "gambar")
OUT = os.path.join(BASE, "Laporan_UTS_Network_Explorer_Kelompok3.docx")
OUT_FALLBACK = os.path.join(BASE, "Laporan_UTS_Network_Explorer_Kelompok3_4Anggota.docx")

FONT = "Times New Roman"
PEACH = "F4B183"
PEACH_LIGHT = "F8CBAD"
YELLOW = "FFD966"
BLUE = "B5C7E8"

doc = Document()

# --- page setup A4, margin 2.54 cm ---
for s in doc.sections:
    s.page_width = Cm(21.0)
    s.page_height = Cm(29.7)
    s.top_margin = Cm(2.54)
    s.bottom_margin = Cm(2.54)
    s.left_margin = Cm(2.54)
    s.right_margin = Cm(2.54)

# --- base style ---
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
pf = st.paragraph_format
pf.line_spacing = 1.5
pf.space_after = Pt(6)
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def para(text, bold=False, italic=False, align="justify", size=12, space_after=6, caps=False):
    p = doc.add_paragraph()
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    if caps:
        text = text.upper()
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    return p


def para_mixed(parts, align="justify"):
    """parts: list of (text, bold, italic)"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if align == "justify" else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    for t, b, i in parts:
        r = p.add_run(t)
        r.font.name = FONT
        r.font.size = Pt(12)
        r.bold = b
        r.italic = i
    return p


def heading_letter(letter, title):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(6)
    r1 = p.add_run(letter + ".  " + title)
    r1.font.name = FONT
    r1.font.size = Pt(12)
    r1.bold = True
    return p


def sub_number(num, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(str(num) + ".  " + title)
    r.font.name = FONT
    r.font.size = Pt(12)
    r.bold = True
    return p


def numbered(n, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(str(n) + ".  " + text)
    r.font.name = FONT
    r.font.size = Pt(12)
    return p


def lettered(code, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(code + ".  " + text)
    r.font.name = FONT
    r.font.size = Pt(12)
    return p


def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.clear()
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.font.name = FONT
        r.font.size = Pt(12)
        r.bold = True
        r2 = p.add_run(text)
        r2.font.name = FONT
        r2.font.size = Pt(12)
    else:
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(12)
    return p


def make_table(headers, rows, widths=None, header_color=None, header_bold=True):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        c = t.cell(0, j)
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.name = FONT
        r.font.size = Pt(12)
        r.bold = header_bold
        if header_color:
            shade(c, header_color)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = FONT
            r.font.size = Pt(11)
    if widths:
        for j, w in enumerate(widths):
            for i in range(len(rows) + 1):
                t.cell(i, j).width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return t


def add_image(name, caption):
    path = os.path.join(IMG, name)
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        r.add_picture(path, width=Cm(14.0))
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(8)
        rr = c.add_run(caption)
        rr.font.name = FONT
        rr.font.size = Pt(11)
        rr.italic = True
    else:
        para("[Gambar " + name + " belum tersedia]", italic=True, align="center")


# ================= COVER =================
para("UJIAN TENGAH SEMESTER (UTS)", bold=True, align="center", size=14)
para("PENGEMBANGAN PROTOTYPE MULTIMEDIA", bold=True, align="center", size=12)
para("PEMBELAJARAN INTERAKTIF", bold=True, align="center", size=12)
para("Disusun Guna Memenuhi Tugas Mata TIK dalam Pendidikan", italic=True, align="center")
para("", size=12)
para("Dosen Pengampu:", bold=True, align="center")
para("Reni Rahmadani, S.Kom., M.Kom.", align="center")
para("Prof.Dr. Sahat Siagian, M.Pd", align="center")
para("", size=12)
lp = doc.add_paragraph()
lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
lr = lp.add_run()
lr.add_picture(os.path.join(IMG, "logo-unimed.png"), width=Cm(4.5))
para("", size=12)
para("Disusun Oleh:", bold=True, align="center")
para("Kelompok 3", bold=True, align="center")

t = doc.add_table(rows=5, cols=3)
t.style = "Table Grid"
cover_rows = [("Nama", ":", "NIM"), ("Anggota 1 (Ketua)", ":", "525XXXXX01"),
              ("Anggota 2", ":", "525XXXXX02"), ("Anggota 3", ":", "525XXXXX03"),
              ("Anggota 4", ":", "525XXXXX04")]
for i, (a, b, c) in enumerate(cover_rows):
    for j, v in enumerate((a, b, c)):
        cell = t.cell(i, j)
        cell.text = ""
        r = cell.paragraphs[0].add_run(v)
        r.font.name = FONT
        r.font.size = Pt(12)
        if i == 0:
            r.bold = True
doc.add_paragraph()
para("PROGRAM STUDI PENDIDIKAN TEKNOLOGI INFORMATIKA DAN KOMPUTER", bold=True, align="center")
para("FAKULTAS TEKNIK", bold=True, align="center")
para("UNIVERSITAS NEGERI MEDAN", bold=True, align="center")
para("2026", bold=True, align="center")
doc.add_page_break()

# ================= JUDUL ISI =================
para("PENGENALAN TOPOLOGI JARINGAN BUS, RING, DAN STAR", bold=True, align="center")

# ================= A =================
heading_letter("A", "Identitas Proyek")
t = doc.add_table(rows=9, cols=3)
t.style = "Table Grid"
ident = [
    ("Komponen", ":", "Keterangan"),
    ("Nama proyek", ":", "Prototype Multimedia Pembelajaran Interaktif Network Explorer: Misi Topologi"),
    ("Mata pelajaran", ":", "Informatika"),
    ("Materi", ":", "Pengenalan Topologi Jaringan Bus, Ring, dan Star"),
    ("Target peserta didik", ":", "SMA Kelas XI"),
    ("Media digital", ":", "Website Interaktif (Next.js, diakses lewat peramban)"),
    ("Tautan prototype", ":", "https://network-explorer-xi.vercel.app/"),
    ("Alokasi waktu", ":", "50 menit"),
    ("Bentuk produk", ":", "Prototype multimedia pembelajaran digital berbasis web"),
]
for i, (a, b, c) in enumerate(ident):
    for j, v in enumerate((a, b, c)):
        cell = t.cell(i, j)
        cell.text = ""
        r = cell.paragraphs[0].add_run(v)
        r.font.name = FONT
        r.font.size = Pt(12)
        if i == 0:
            r.bold = True
doc.add_paragraph()

# ================= B =================
heading_letter("B", "Latar Belakang Pengembangan Prototype")
para("Pada Tugas Rutin II, kelompok memilih materi Pengenalan Topologi Jaringan untuk peserta didik SMA kelas XI. Pemilihan materi ini didasarkan pada pentingnya pemahaman jaringan sebagai dasar dalam pembelajaran Informatika. Topologi Bus, Ring, dan Star dipilih karena sering muncul dalam materi kelas XI dan dekat dengan praktik di laboratorium sekolah.")
para("Permasalahan yang ingin diselesaikan adalah masih adanya siswa yang menghafal bentuk topologi tetapi belum memahami bagaimana data bergerak dalam jaringan. Selain itu, siswa juga dapat mengalami kesulitan dalam membedakan dampak gangguan pada tiap topologi, misalnya beda putus backbone pada Bus dengan putus kabel PC pada Star.")
para("Pada storyboard awal, pembelajaran dirancang dalam beberapa layar yang berisi pembuka, tujuan pembelajaran, materi tiga topologi, praktik menyusun jaringan, simulasi pengiriman data, latihan, dan evaluasi. Rancangan tersebut kemudian dikembangkan lagi pada prototype UTS agar memenuhi komponen yang dipersyaratkan, yaitu materi, aktivitas interaktif, assessment, feedback, dan hasil pembelajaran.")

# ================= C =================
heading_letter("C", "Target Peserta Didik")
para("Target pengguna prototype adalah siswa SMA kelas XI yang sedang mempelajari materi Informatika tentang jaringan komputer.")
para("Siswa pada tingkat ini diharapkan sudah dapat menggunakan komputer dan peramban secara umum, tetapi belum semuanya memahami konsep topologi jaringan. Oleh karena itu, materi disampaikan menggunakan bahasa yang sederhana dan contoh yang dekat dengan kegiatan di laboratorium sekolah.")
para_mixed([("Contoh yang digunakan dalam prototype adalah ", False, False),
            ("laboratorium dengan 12 sampai 24 PC", True, False),
            (". Contoh tersebut dipilih karena dekat dengan pengalaman siswa sehingga siswa dapat melihat hubungan antara susunan kabel, aliran data, dan gangguan jaringan.", False, False)])

# ================= D =================
heading_letter("D", "Materi atau Topik Pembelajaran")
para("Materi terdiri dari beberapa bagian, yaitu:")
numbered(1, "Pengertian jaringan komputer dan topologi.")
numbered(2, "Ciri dan cara kerja topologi Bus.")
numbered(3, "Ciri dan cara kerja topologi Ring.")
numbered(4, "Ciri dan cara kerja topologi Star.")
numbered(5, "Perbandingan Bus, Ring, dan Star.")
numbered(6, "Praktik menyusun jaringan Star.")
numbered(7, "Simulasi pengiriman data dan gangguan jaringan.")
numbered(8, "Evaluasi pemahaman siswa.")
para("Pada bagian materi, siswa diberikan contoh laboratorium sekolah agar konsep topologi dapat dipahami melalui keadaan nyata. Selanjutnya, contoh tersebut dihubungkan dengan simulasi pengiriman paket data sehingga siswa dapat memahami bahwa pilihan topologi menentukan jalur data dan dampak gangguan secara visual.")

# ================= E =================
heading_letter("E", "Masalah atau Kesulitan Belajar yang Ingin di Selesaikan")
para("Kesulitan utama yang ingin diselesaikan melalui prototype ini adalah siswa masih dapat mengalami kebingungan dalam membedakan topologi Bus, Ring, dan Star.")
para("Topologi Bus lebih mudah dipahami sebagai satu kabel utama yang dipakai bersama, sedangkan Ring menggunakan jalur melingkar dan Star menggunakan switch pusat. Jika ketiga konsep tersebut hanya dijelaskan menggunakan teks dan gambar diam, siswa dapat kesulitan memahami aliran data dan hubungan antara bentuk jaringan dengan dampak gangguannya.")
para("Selain itu, siswa juga perlu memahami fungsi dari masing-masing perangkat seperti backbone, terminator, dan switch. Oleh karena itu, prototype dibuat dengan memberikan diagram animasi, praktik menyusun kabel, simulasi ping, soal evaluasi, serta feedback yang menjelaskan alasan jawaban.")
para("Dengan adanya interaksi tersebut, siswa tidak hanya membaca materi, tetapi juga dapat mencoba menjawab dan melihat kembali bagian yang masih belum dipahami.")

# ================= F =================
heading_letter("F", "Tujuan Pembelajaran")
para("Setelah menggunakan prototype multimedia pembelajaran ini, peserta didik diharapkan mampu:")
lettered("a", "Menjelaskan pengertian topologi jaringan dengan bahasa sederhana.")
lettered("b", "Menjelaskan cara kerja topologi Bus, Ring, dan Star.")
lettered("c", "Mengenali dan menjelaskan fungsi backbone, terminator, dan switch.")
lettered("d", "Menjelaskan dampak gangguan pada tiap topologi.")
lettered("e", "Memilih topologi yang sesuai untuk kebutuhan laboratorium sekolah.")
para("Tujuan tersebut tetap mempertahankan tujuan pembelajaran dari rancangan awal sehingga pengembangan prototype tidak mengubah fokus utama materi.")

# ================= G =================
heading_letter("G", "Perubahan dan Penyempurnaan dari Rancangan Awal")
para("Pengembangan prototype UTS tidak mengubah topik utama dari Tugas Rutin II. Perubahan dilakukan pada bagian struktur, interaksi, assessment, feedback, dan hasil agar prototype dapat memenuhi ketentuan UTS.")
make_table(
    ["Bagian", "Rancangan Awal", "Perubahan pada Prototype", "Alasan Perubahan"],
    [
        ["Pembuka", "Halaman pembuka berisi judul materi dan identitas", "Dibuat halaman pembuka dengan judul Misi Topologi, identitas kelas XI, dan tombol Mulai Misi", "Agar siswa mengetahui materi yang akan dipelajari dan cara memulai pembelajaran."],
        ["Materi topologi", "Materi tiga topologi masih dalam bentuk teks", "Materi dibuat lebih bertahap dan dilengkapi diagram animasi, narasi, dan contoh lab sekolah", "Agar konsep aliran data lebih mudah dipahami melalui tampilan visual."],
        ["Perbandingan", "Belum ada perbandingan antar topologi", "Ditambahkan halaman perbandingan 7 aspek dan studi kasus lab 24 PC", "Agar siswa tidak hanya melihat bentuk topologi tetapi memahami penggunaannya."],
        ["Praktik", "Hanya ada contoh gambar jaringan", "Menambahkan praktik susun kabel Star dengan drag and drop", "Agar siswa lebih mudah memahami penerapan perangkat sebelum mengerjakan evaluasi."],
        ["Simulasi", "Belum ada simulasi gangguan", "Ditampilkan simulasi kirim data, putus kabel PC-03, dan penanda kesehatan koneksi", "Agar siswa dapat melihat hubungan antara struktur jaringan dan dampak gangguan."],
        ["Evaluasi", "Evaluasi awal terdiri dari tiga soal", "Evaluasi dikembangkan menjadi 10 soal", "Untuk memenuhi ketentuan UTS yang meminta minimal lima butir assessment."],
        ["Feedback", "Feedback awal hanya benar atau salah", "Feedback tidak hanya menunjukkan benar atau salah, tetapi memberikan penjelasan", "Agar siswa mengetahui alasan jawaban dan dapat memahami kesalahannya."],
        ["Hasil", "Belum ada halaman hasil", "Ditambahkan halaman hasil dan rangkuman berisi skor, XP, lencana, dan rekomendasi", "Agar siswa mengetahui hasil pembelajaran serta bagian yang perlu dipelajari kembali."],
        ["Navigasi", "Alur halaman belum jelas", "Menggunakan navbar, menu misi, dan tombol antarhalaman", "Agar siswa dapat berpindah halaman sesuai alur pembelajaran."],
    ],
    widths=[2.5, 3.8, 4.8, 4.8],
    header_color=PEACH,
)

# ================= H =================
heading_letter("H", "Rancangan Prototype Multimedia")
para("Prototype dibuat menggunakan website interaktif berbasis Next.js yang diakses lewat peramban dan dipublikasi di Vercel. Website dipilih karena sesuai dengan rancangan awal dan dapat digunakan untuk menggabungkan teks, diagram animasi, simulasi, tombol, navigasi, contoh, evaluasi, dan feedback dalam satu alur.")
para("Prototype yang telah dibuat terdiri dari halaman pembuka, misi utama, materi Bus Ring Star, praktik menyusun kabel, sandbox simulasi, perbandingan, evaluasi, feedback, hasil, dan progres belajar.")
para("Prototype dapat diakses pada tautan https://network-explorer-xi.vercel.app/ melalui peramban pada komputer atau ponsel. Untuk membuka halaman misi, materi, praktik, dan kuis, gunakan akun demo navigator@example.com dengan kata sandi password123 pada halaman masuk.")
para("Materi pada prototype menggunakan minimal tiga jenis media, yaitu teks penjelasan, diagram animasi SVG, dan simulasi interaktif, sehingga memenuhi ketentuan minimal dua jenis media.")
sub_number(1, "Struktur Halaman Prototype")
make_table(
    ["Bagian", "Isi", "Interaksi Pengguna"],
    [
        ["Pembuka", "Judul, materi, dan tombol Mulai Misi", "Menekan tombol Mulai Misi"],
        ["Misi Utama", "Pilihan Level 1 Bus, Level 2 Ring, Level 3 Star", "Memilih level dan membuka materi"],
        ["Materi Bus", "Pengertian, cara kerja, kelebihan, dan kekurangan Bus", "Mengikuti langkah materi"],
        ["Materi Ring", "Pengertian, cara kerja, kelebihan, dan kekurangan Ring", "Mengikuti langkah materi"],
        ["Materi Star", "Pengertian, cara kerja, kelebihan, dan kekurangan Star", "Memilih diagram dan cek pemahaman"],
        ["Praktik", "Menyusun PC, switch, dan kabel dengan drag and drop", "Menyambung kabel ke slot"],
        ["Sandbox", "Simulasi kirim data dan putus kabel PC-03", "Menekan kirim data dan mengamati log"],
        ["Perbandingan", "Tabel 7 aspek dan studi kasus lab 24 PC", "Membaca dan membandingkan"],
        ["Evaluasi", "Sepuluh soal pilihan ganda", "Menjawab soal"],
        ["Feedback", "Penjelasan jawaban benar atau salah", "Melihat penjelasan dan mencoba kembali"],
        ["Hasil dan Rangkuman", "Skor, XP, lencana, dan rekomendasi", "Mengulang materi atau evaluasi atau selesai"],
    ],
    widths=[2.8, 6.5, 6.5],
    header_color=PEACH_LIGHT,
)
para("Struktur tersebut dikembangkan dari storyboard awal. Pada prototype, bagian yang sebelumnya masih berupa rancangan layar diwujudkan menjadi halaman website yang memiliki tombol dan navigasi.")
sub_number(2, "Alur Navigasi")
para("Alur utama prototype adalah:")
para("Pembuka → Misi Utama → Materi Bus → Materi Ring → Materi Star → Praktik → Sandbox → Evaluasi → Feedback → Hasil.", align="left")
para("Pada bagian evaluasi, pengguna dapat memilih jawaban. Jika jawaban benar, pengguna mendapatkan penjelasan dan dapat melanjutkan. Jika jawaban salah, pengguna diarahkan ke feedback yang menjelaskan jawaban yang dipilih, jawaban yang benar, serta alasan jawaban tersebut.")
para("Pada halaman hasil, siswa dapat melihat hasil evaluasi dan memilih untuk mengulang materi atau menyelesaikan pembelajaran.")

# ================= I =================
heading_letter("I", "Tampilan Prototype")
para("Prototype multimedia pembelajaran interaktif dibuat menggunakan website dengan tampilan yang sederhana dan mudah digunakan oleh siswa kelas XI. Setiap halaman dilengkapi dengan navigasi dan interaksi sesuai dengan alur pembelajaran.")
para("Tampilan prototype terdiri dari:")
numbered(1, "Tampilan Halaman Pembuka")
add_image("01-pembuka.png", "Gambar 1. Halaman pembuka Network Explorer: Misi Topologi.")
numbered(2, "Tampilan Misi Utama")
add_image("03-misi.png", "Gambar 2. Halaman misi utama berisi pilihan Level 1 Bus, Level 2 Ring, dan Level 3 Star.")
numbered(3, "Tampilan Materi Star")
add_image("04-materi-star.png", "Gambar 3. Halaman materi Star dengan diagram animasi dan cek pemahaman.")
numbered(4, "Tampilan Praktik Menyusun Kabel")
add_image("05-praktik-star.png", "Gambar 4. Halaman praktik menyusun kabel Star dengan drag and drop.")
numbered(5, "Tampilan Sandbox Simulasi")
add_image("06-sandbox.png", "Gambar 5. Halaman sandbox untuk uji kirim data dan simulasi kabel putus.")
numbered(6, "Tampilan Kuis dan Hasil")
add_image("07-kuis.png", "Gambar 6. Halaman kuis akhir dengan navigasi nomor soal dan progres pengerjaan.")
add_image("09-hasil.png", "Gambar 7. Halaman hasil berisi skor, XP, lencana, dan pembahasan soal.")

# ================= J =================
heading_letter("J", "Aktivitas Interaktif")
para("Aktivitas interaktif dibuat agar siswa tidak hanya membaca materi, tetapi juga melakukan tindakan selama pembelajaran.")
para("Beberapa bentuk interaksi dalam prototype yaitu:")
numbered(1, "Menekan tombol Mulai Misi untuk memulai pembelajaran.")
numbered(2, "Memilih level Bus, Ring, atau Star pada halaman misi.")
numbered(3, "Menggunakan tombol navigasi untuk berpindah bagian.")
numbered(4, "Melihat diagram animasi aliran data tiap topologi.")
numbered(5, "Menyusun PC, switch, dan kabel dengan drag and drop.")
numbered(6, "Menekan tombol kirim data dan memutus kabel pada sandbox.")
numbered(7, "Menjawab sepuluh soal evaluasi.")
numbered(8, "Melihat feedback dari jawaban.")
numbered(9, "Melihat hasil, XP, dan rangkuman serta mengulang materi jika diperlukan.")
para("Aktivitas tersebut membuat siswa memiliki peran dalam menggunakan media. Siswa tidak hanya menjadi pembaca materi, tetapi juga harus memilih jawaban dan mengikuti alur pembelajaran.")

# ================= K =================
heading_letter("K", "Assessment dan Feedback")
bullet("Evaluasi", bold_prefix="Evaluasi")
para("Pada rancangan Tugas Rutin II, evaluasi awalnya terdiri dari tiga soal. Pada prototype UTS, jumlah evaluasi dikembangkan menjadi sepuluh soal agar mencakup Bus, Ring, Star, perangkat, dan studi kasus lab.")
para("Berikut ringkasan 10 soal yang ada pada website:")
para("Kesepuluh soal disusun sesuai tujuan pembelajaran: soal tentang pengertian untuk tujuan a, soal cara kerja Bus Ring Star untuk tujuan b, soal backbone terminator dan switch untuk tujuan c, soal dampak gangguan dan ping untuk tujuan d, serta soal studi kasus lab untuk tujuan e.")
soal = [
    ("1. Apa yang terjadi jika switch pusat gagal? (Studi kasus lab 12 PC Star)", "B", "Switch adalah titik pusat Star. Jika mati, tidak ada penerus frame sehingga seluruh komunikasi via switch berhenti."),
    ("2. Mengapa satu kabel putus melumpuhkan Ring tunggal? (5 PC searah jarum jam)", "B", "Ring tunggal hanya punya satu jalur melingkar. Satu titik putus membuat lingkaran terbuka sehingga paket tidak sampai."),
    ("3. Lab sekolah butuh 24 PC dan satu gangguan tidak boleh melumpuhkan lab. Topologi yang tepat?", "B", "Topologi Star memakai switch sebagai pusat. Jika satu kabel PC rusak, PC lain masih bisa berkomunikasi."),
    ("4. Apa fungsi terminator pada ujung backbone Bus?", "B", "Terminator menyerap sinyal di ujung backbone agar tidak terjadi pantulan yang merusak paket."),
    ("5. Kapan collision paling rawan terjadi?", "B", "Media bersama pada Bus rentan collision saat ramai. Star dengan switch dan Ring dengan token menghindarinya."),
    ("6. Bagaimana switch meneruskan frame dari PC-01 untuk PC-04?", "B", "Switch belajar alamat MAC dan meneruskan unicast yang dikenal hanya ke port tujuan."),
    ("7. Ping dari PC-01 ke PC-03 gagal total. Langkah pertama yang tepat?", "B", "Urutkan pemeriksaan dari fisik ke logis: kabel dan switch, lalu alamat IP, lalu firewall yang mungkin memblokir ICMP."),
    ("8. Apa yang dimaksud backbone pada Bus klasik?", "A", "Backbone adalah kabel utama bersama yang dipakai semua node pada Bus klasik."),
    ("9. Mengapa Star butuh lebih banyak kabel dibanding Bus untuk 12 PC?", "B", "Star memakai satu kabel per node ke switch, sedangkan Bus berbagi satu backbone sehingga lebih hemat."),
    ("10. Apa keunggulan ring ganda dibanding ring tunggal?", "B", "Ring ganda memberi jalur cadangan. Saat satu jalur putus, trafik dialihkan ke jalur cadangan."),
]
for q, jwb, fb in soal:
    para(q, align="left")
    para("Jawaban: " + jwb, align="left")
    para("Feedback: " + fb)
bullet("Kriteria Hasil", bold_prefix="Kriteria Hasil")
make_table(
    ["Rentang Skor", "Keterangan", "Tindak Lanjut"],
    [
        ["80-100", "Sangat baik", "Siswa dapat melanjutkan atau mencoba latihan tambahan."],
        ["60-79", "Cukup", "Siswa disarankan mengulang bagian perbandingan dan simulasi Star."],
        ["0-59", "Perlu belajar kembali", "Siswa diarahkan mempelajari kembali materi dasar kemudian mengulang evaluasi."],
    ],
    widths=[2.5, 3.0, 10.4],
)

# ================= L =================
heading_letter("L", "Penerapan Prinsip Desain Multimedia")
sub_number(1, "Segmenting")
para("Materi dibagi menjadi beberapa bagian, yaitu pengenalan jaringan, materi Bus, materi Ring, materi Star, perbandingan, praktik, simulasi, evaluasi, dan hasil.")
para("Pembagian tersebut membuat materi tidak diberikan sekaligus dalam satu halaman sehingga siswa dapat mempelajari materi secara bertahap.")
sub_number(2, "Signaling")
para("Bagian penting seperti nama topologi, kata kunci, nomor soal, penanda switch pusat, dan bagian feedback diberi penekanan.")
para("Penerapan ini bertujuan membantu siswa menemukan informasi penting dengan lebih mudah.")
sub_number(3, "Coherence")
para("Isi prototype dibuat berfokus pada materi topologi Bus, Ring, dan Star. Informasi yang tidak berhubungan langsung dengan pembelajaran tidak dimasukkan.")
para("Hal tersebut dilakukan agar siswa dapat lebih fokus terhadap materi yang sedang dipelajari.")
sub_number(4, "Contiguity")
para("Penjelasan mengenai perangkat dan aliran data ditempatkan berdekatan dengan diagram yang dijelaskan. Dengan cara tersebut, siswa dapat lebih mudah menghubungkan antara informasi berupa teks dengan bentuk diagram.")
para("Penggunaan empat prinsip tersebut membantu membuat prototype lebih terstruktur dan tidak hanya berisi teks.")

# ================= M =================
heading_letter("M", "Integrasi Teknologi dalam Pembelajaran (TPACK)")
para("Pengembangan prototype dapat dijelaskan menggunakan kerangka TPACK, yaitu hubungan antara pengetahuan materi, pengetahuan pedagogis, dan pengetahuan teknologi.")
make_table(
    ["Komponen", "Penerapan"],
    [
        ["Content Knowledge (CK)", "Pengetahuan tentang konsep jaringan, cara kerja Bus Ring Star, perangkat backbone terminator switch, dan dampak gangguan."],
        ["Pedagogical Knowledge (PK)", "Pembelajaran bertahap melalui penjelasan, contoh lab sekolah, aktivitas praktik, simulasi, evaluasi, dan feedback."],
        ["Technology Knowledge (TK)", "Penggunaan website Next.js, navigasi, diagram SVG animasi, drag and drop, simulasi ping, log kejadian, dan sistem XP."],
        ["TPACK", "Teknologi digunakan untuk menyajikan materi Informatika dengan cara yang lebih visual dan interaktif, bukan hanya sebagai tempat menampilkan teks."],
    ],
    widths=[4.5, 11.4],
    header_color=YELLOW,
)

# ================= N =================
heading_letter("N", "Posisi Teknologi Berdasarkan SAMR")
para("Berdasarkan model SAMR, penggunaan teknologi dalam prototype dapat dilihat dari beberapa tingkatan.")
bullet("Materi pembelajaran yang sebelumnya dapat disampaikan melalui media cetak disajikan dalam bentuk digital menggunakan website.", bold_prefix="Substitution")
bullet("Website diberikan tambahan tombol navigasi, diagram animasi, dan feedback sehingga penggunaannya menjadi lebih interaktif dibandingkan penyajian materi biasa.", bold_prefix="Augmentation")
bullet("Prototype mengubah rancangan pembelajaran menjadi lebih interaktif melalui praktik susun kabel, simulasi ping, evaluasi, feedback, dan halaman hasil.", bold_prefix="Modification")
bullet("Simulasi gangguan kabel dan pengiriman paket secara langsung memberi pengalaman baru yang sulit dilakukan tanpa teknologi web.", bold_prefix="Redefinition")
para("Dalam prototype ini, penggunaan teknologi terutama berada pada tahap Modification menuju Redefinition, karena website tidak hanya menggantikan media penyampaian materi, tetapi digunakan untuk mengubah alur pembelajaran menjadi pengalaman yang melibatkan tindakan pengguna.")

# ================= O =================
heading_letter("O", "Pembagian Pekerjaan Kelompok")
make_table(
    ["Anggota", "Bagian yang Dikerjakan", "Bentuk Kontribusi"],
    [
        ["Anggota 1 (Ketua tim)", "Materi dan pengembangan storyboard", "Menyusun materi, tujuan pembelajaran, alur materi, dan rancangan interaksi. Ketua tim memastikan seluruh komponen prototype terintegrasi."],
        ["Anggota 2", "Desain visual", "Membuat tampilan halaman, tata letak, diagram, dan elemen visual website."],
        ["Anggota 3", "Interaksi dan navigasi", "Membuat alur halaman, tombol, navigasi, praktik drag and drop, dan simulasi sandbox."],
        ["Anggota 4", "Assessment dan pengujian", "Menyusun soal, feedback, dan melakukan pengujian prototype."],
    ],
    widths=[3.0, 4.5, 8.4],
    header_color=BLUE,
)
para("Catatan: nama Anggota 1 sampai Anggota 4 masih berupa placeholder dan akan diganti dengan nama dan NIM asli.")
para("Seluruh anggota memahami produk yang dikembangkan. AI hanya digunakan sebagai alat bantu penulisan dan screenshot, sedangkan kelompok tetap bertanggung jawab atas isi dan produk akhir.")

# ================= P =================
heading_letter("P", "Kesimpulan")
para("Prototype multimedia pembelajaran interaktif Network Explorer: Misi Topologi dikembangkan berdasarkan storyboard yang telah dibuat pada Tugas Rutin II. Topik, target peserta didik, dan tujuan pembelajaran tetap dipertahankan, sedangkan beberapa bagian dikembangkan agar prototype dapat memenuhi ketentuan UTS.")
para("Pengembangan dilakukan dengan membuat materi lebih terstruktur, menambahkan aktivitas interaktif berupa praktik susun kabel dan simulasi sandbox, menggunakan navbar dan tombol sebagai navigasi, serta mengembangkan evaluasi dari tiga soal menjadi sepuluh soal. Feedback juga dibuat lebih lengkap karena tidak hanya menunjukkan benar atau salah, tetapi memberikan penjelasan mengenai jawaban yang dipilih. Selain itu, ditambahkan halaman hasil dan rangkuman sebagai bagian penutup pembelajaran.")
para("Prototype juga menerapkan prinsip desain multimedia berupa Segmenting, Signaling, Coherence, dan Contiguity. Integrasi teknologi dijelaskan melalui TPACK, sedangkan posisi teknologi dalam pembelajaran dianalisis menggunakan SAMR dan terutama berada pada tahap Modification menuju Redefinition.")
para("Secara keseluruhan, prototype diharapkan dapat membantu siswa kelas XI memahami konsep topologi Bus, Ring, dan Star dengan cara yang lebih sederhana dan interaktif. Sebelum digunakan untuk demonstrasi dan dikumpulkan, prototype tetap perlu diuji kembali untuk memastikan semua tombol, navigasi, soal, feedback, dan simulasi dapat berjalan dengan baik.")

# ================= DAFTAR SUMBER =================
para("DAFTAR SUMBER DAN ASET", bold=True, align="center")
para("Mawarni, S., & Muhtadi, A. (2017). Pengembangan digital book interaktif mata kuliah pengembangan multimedia pembelajaran interaktif untuk mahasiswa teknologi pendidikan. Jurnal Inovasi Teknologi Pendidikan, 4(1), 84-96.")
para("Saluky, & Riyanto, O. R. (2025). Multimedia pembelajaran. CV. Zenius Publisher.")
para("Mayer, R. E. (2021). Multimedia learning (3rd ed.). Cambridge University Press.")
para("Aset/Media:", bold=True, align="left")
para("Diagram topologi, ikon perangkat, dan elemen visual pada prototype dibuat dengan kode SVG pada website Network Explorer. Screenshot tampilan diambil langsung dari https://network-explorer-xi.vercel.app/ sebagai dokumentasi milik kelompok. Website dibangun dengan Next.js dan dipublikasi melalui Vercel.")

try:
    doc.save(OUT)
    print("saved:", OUT)
except PermissionError:
    doc.save(OUT_FALLBACK)
    print("utama terkunci, saved:", OUT_FALLBACK)
